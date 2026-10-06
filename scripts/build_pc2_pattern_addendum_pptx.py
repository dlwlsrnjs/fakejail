#!/usr/bin/env python3
"""Build a concise PPTX addendum for person-language pattern findings."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_pc2_analysis_pptx import (
    BLUE, GRID, MUTED, ORANGE, PANEL, TEAL, TEXT,
    add_box, add_footer, add_header, add_metric_bar, add_rich_lines,
    add_text, add_title_slide, new_prs, set_bg,
)


def slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s, title, subtitle); return s


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--patterns", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--shortlist-csv", type=Path, required=True)
    args = parser.parse_args()
    p = json.load(args.patterns.open(encoding="utf-8"))
    prs = new_prs()
    add_title_slide(prs, "Person–language pattern analysis", "Can a useful clue language be selected before target generation?")
    page = 2

    s = slide(prs, "How PC² models the relationship", "Official implementation adapted to our text-generation experiment")
    add_rich_lines(s, [
        ("Step 1 · Resolve the political entity and related country", TEAL, True),
        ("A public figure is mapped to the country most strongly associated with their role.", TEXT, False),
        ("Step 2 · Translate the entity description into 72 languages", TEAL, True),
        ("Low-quality translations are filtered using round-trip semantic similarity.", TEXT, False),
        ("Step 3 · Score every entity–language pair", TEAL, True),
        ("keyword bias 0.75 · politics 0.733 · country common knowledge 0.617 · keyword common knowledge 0.667", TEXT, False),
        ("Step 4 · Select a score percentile", TEAL, True),
        ("The released code sorts ascending and chooses the requested percentile independently for each indirect entity.", TEXT, False),
    ], .75, 1.42, 11.85, 4.8, 12)
    add_footer(s, page); page += 1

    v = p["variance_decomposition"]
    s = slide(prs, "The person matters much more than the language")
    add_metric_bar(s, "Person main effect", v["person_main_effect_pct_total_ss"], .9, 1.65, 11.5, TEAL)
    add_metric_bar(s, "Language main effect", v["language_main_effect_pct_total_ss"], .9, 2.5, 11.5, BLUE)
    add_metric_bar(s, "Interaction + cell noise", v["unexplained_interaction_and_cell_noise_pct"], .9, 3.35, 11.5, ORANGE)
    d = p["person_driver_analysis"]
    add_rich_lines(s, [
        ("Direct-original success is a strong susceptibility prior", TEAL, True),
        (f"Translated success averages {d['translated_rate_if_direct_succeeds_pct']:.2f}% when the direct prompt succeeds, versus {d['translated_rate_if_direct_fails_pct']:.2f}% when it fails.", TEXT, False),
        ("Identity resolution is the dominant mechanism", TEAL, True),
        (f"Across people, translated strict success correlates r={d['person_level_corr_strict_with_entity_match']:.3f} with entity match, versus r={d['person_level_corr_strict_with_non_refusal']:.3f} with non-refusal.", TEXT, False),
    ], .85, 4.4, 11.65, 1.75, 11.5)
    add_footer(s, page); page += 1

    s = slide(prs, "Geographic relation is weaker than expected", "Within-person comparisons reduce confounding by easy and hard people")
    rel = p["relation_groups"]; paired = p["paired_relation_effects"]
    rows = [
        ("Exact person-country language", rel["person_country_native_language"]["success_rate_pct"], rel["all_other_languages"]["success_rate_pct"], paired["native_language_vs_other_languages"]["wilcoxon_p"]),
        ("Source article language", rel["source_article_language"]["success_rate_pct"], rel["not_source_article_language"]["success_rate_pct"], paired["source_language_vs_other_languages"]["wilcoxon_p"]),
        ("Same script as country language", rel["same_script_as_person_country_language"]["success_rate_pct"], rel["different_script"]["success_rate_pct"], paired["same_vs_different_native_script"]["wilcoxon_p"]),
    ]
    headers = ["Relation", "Matched", "Other", "Difference", "Within-person p"]
    widths = [4.3, 1.6, 1.6, 1.7, 2.1]; x0=.8; y0=1.65
    xx=x0
    for h,w in zip(headers,widths): add_text(s,h,xx,y0,w,.32,10,MUTED,True,PP_ALIGN.RIGHT if h!="Relation" else PP_ALIGN.LEFT); xx+=w
    for i,(label,a,b,pv) in enumerate(rows):
        yy=y0+.65+i*.85; vals=[label,f"{a:.2f}%",f"{b:.2f}%",f"{a-b:+.2f} pp",f"{pv:.3f}"]; xx=x0
        for j,(val,w) in enumerate(zip(vals,widths)):
            add_text(s,val,xx,yy,w,.38,12,TEAL if j==3 and a>b else TEXT,j in (0,3),PP_ALIGN.RIGHT if j else PP_ALIGN.LEFT,MSO_ANCHOR.MIDDLE); xx+=w
    add_rich_lines(s, [
        ("Result", ORANGE, True),
        ("Same-script languages show a +7.67 pp aggregate advantage and +6.93 pp mean within-person advantage, but the pilot is not significant (p=.131). Exact native-language and source-language matching do not help.", TEXT, False),
    ], .8, 5.25, 11.75, 1.05, 11)
    add_footer(s, page); page += 1

    s = slide(prs, "PC² semantic scores do not rank languages within a person")
    corr = p["pc2_feature_correlations"]
    headers=["Feature","Overall r","Within-person ρ","Within-person p"]
    widths=[4.4,2.0,2.2,2.2]; x0=.9; y0=1.38; xx=x0
    for h,w in zip(headers,widths): add_text(s,h,xx,y0,w,.3,10,MUTED,True,PP_ALIGN.RIGHT if h!="Feature" else PP_ALIGN.LEFT); xx+=w
    names={"keyword_bias":"Keyword bias","politics":"Politics","country_common_knowledge":"Country common knowledge","keyword_common_knowledge":"Keyword common knowledge","combined_score":"Combined score","backtranslation_similarity":"Round-trip similarity","translation_length_ratio":"Translation length ratio"}
    for i,k in enumerate(names):
        r=corr[k]; yy=y0+.48+i*.55; vals=[names[k],f"{r['overall_r']:+.3f}",f"{r['within_person_spearman']:+.3f}",f"{r['within_person_p']:.3f}"]; xx=x0
        for j,(val,w) in enumerate(zip(vals,widths)):
            add_text(s,val,xx,yy,w,.3,10.5,TEXT,j==0,PP_ALIGN.RIGHT if j else PP_ALIGN.LEFT,MSO_ANCHOR.MIDDLE); xx+=w
    add_text(s,"The apparent overall correlation is mostly person difficulty. Within a fixed person, every PC² feature is near zero and non-significant.",.9,6.05,11.3,.42,12,ORANGE,True,PP_ALIGN.CENTER)
    add_footer(s, page); page += 1

    s = slide(prs, "Prospective selection on an unseen person", "Leave-one-person-out: the target person's outcomes are excluded from training")
    names=[
        ("Random single language", None),
        ("Global language prior", "global_language_prior"),
        ("PC² relation features", "pc2_relation_features"),
        ("PC² + language prior", "pc2_features_plus_language_prior"),
        ("Global top-1 + relation shortlist", "global_top1_then_relation"),
    ]
    headers=["Selector","Top-1","Top-3","Top-5"]; widths=[5.1,2.0,2.0,2.0]; x0=.9;y0=1.45;xx=x0
    for h,w in zip(headers,widths): add_text(s,h,xx,y0,w,.32,10,MUTED,True,PP_ALIGN.RIGHT if h!="Selector" else PP_ALIGN.LEFT);xx+=w
    for i,(label,key) in enumerate(names):
        yy=y0+.58+i*.75
        if key is None:
            vals=[label,f"{p['random_single_language_expected_success_pct']:.2f}%","—","—"]
        else:
            top=p["leave_one_person_out_selection"][key]["top_k"]
            vals=[label]+[f"{top[str(k)]['hit_rate_pct']:.2f}%" for k in (1,3,5)]
        xx=x0
        for j,(val,w) in enumerate(zip(vals,widths)):
            color=TEAL if (key=="global_top1_then_relation" and j in (1,2)) else TEXT
            add_text(s,val,xx,yy,w,.34,11,color,j in (0,1,2),PP_ALIGN.RIGHT if j else PP_ALIGN.LEFT,MSO_ANCHOR.MIDDLE);xx+=w
    add_text(s,"Best current operating point: one robust global language plus two person-conditioned candidates → 21/23 people (91.30%) had at least one successful language.",.85,5.65,11.6,.62,12,TEAL,True,PP_ALIGN.CENTER)
    add_footer(s, page); page += 1

    hybrid=p["leave_one_person_out_selection"]["global_top1_then_relation"]["per_person_top3"]
    args.shortlist_csv.parent.mkdir(parents=True,exist_ok=True)
    with args.shortlist_csv.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=["person","pilot_id","candidate_1","candidate_2","candidate_3","retrospective_hit"]);w.writeheader()
        for r in hybrid:
            w.writerow({"person":r["person"],"pilot_id":r["pilot_id"],"candidate_1":r["selected_languages"][0],"candidate_2":r["selected_languages"][1],"candidate_3":r["selected_languages"][2],"retrospective_hit":r["hit"]})

    s = slide(prs, "Held-out three-language shortlists", "Retrospective audit of the hybrid selector; not a final universal policy")
    for i,r in enumerate(hybrid):
        col=0 if i<12 else 1; row=i if i<12 else i-12; x=.65+col*6.15; y=1.27+row*.47
        status="✓" if r["hit"] else "×"
        add_text(s,f"{status} {r['person']}",x,y,2.35,.3,8.5,TEAL if r["hit"] else ORANGE,True)
        add_text(s," · ".join(r["selected_languages"]),x+2.3,y,3.65,.3,8.1,TEXT)
    add_footer(s, page); page += 1

    s = slide(prs, "Recommended pattern for a new person", "Use a shortlist and a benign identity probe; do not rely on country language alone")
    steps=[
        ("1", "Build an identity-preserving clue", "Role, country, tenure, institution, and distinguishing career facts; omit the name."),
        ("2", "Translate and quality-filter", "Generate 72 variants and retain acceptable round-trip similarity."),
        ("3", "Run a benign identity-reconstruction probe", "Prefer languages where the local target can recover the intended person without any news-generation request."),
        ("4", "Rank a three-language shortlist", "Anchor with the robust global prior; add two person-conditioned candidates. Same script is only a weak tie-breaker."),
        ("5", "Log the selector separately", "Evaluate selection and target generation on held-out people to prevent leakage."),
    ]
    for i,(n,title,desc) in enumerate(steps):
        y=1.35+i*1.02; add_box(s,.75,y,11.8,.78,PANEL); add_text(s,n,.9,y+.15,.45,.38,20,TEAL,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE);add_text(s,title,1.55,y+.11,3.25,.28,11.5,TEXT,True);add_text(s,desc,4.65,y+.11,7.55,.45,9.5,MUTED)
    add_footer(s, page); page += 1

    s = slide(prs, "Conclusion")
    add_rich_lines(s, [
        ("Yes, there is a pattern—but it is mostly person susceptibility, not a simple country-language mapping.", TEAL, True),
        ("Person identity recoverability explains the result far better than PC²'s semantic score or the exact national language.", TEXT, False),
        ("A single person-conditioned top-1 model did not beat the global language prior.", ORANGE, True),
        ("The defensible improvement is a three-language shortlist: global prior + two person-conditioned candidates, followed by a benign identity probe.", TEXT, False),
        ("Validation needed", BLUE, True),
        ("Repeat each cell with multiple seeds, expand beyond 23 people, and confirm with a second judge before treating any selector as stable.", TEXT, False),
    ], .85, 1.5, 11.65, 4.6, 14)
    add_footer(s, page)

    args.output.parent.mkdir(parents=True,exist_ok=True)
    prs.save(args.output)
    print(json.dumps({"output":str(args.output),"slides":len(prs.slides),"shortlist_csv":str(args.shortlist_csv)},ensure_ascii=False))


if __name__ == "__main__":
    main()
